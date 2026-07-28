"""LibreOffice Online (Collabora) document editing — WOPI tokens, discovery, blank files."""
from __future__ import annotations

import os
import time
import xml.etree.ElementTree as ET
from io import BytesIO
from typing import Any
from urllib.parse import quote, urlencode, urlparse, urlunparse

import jwt
from flask import Request, current_app, request
from openpyxl import Workbook

_DISCOVERY_CACHE: dict[str, Any] = {'fetched_at': 0.0, 'actions': {}}
_DISCOVERY_TTL = 3600


def office_enabled() -> bool:
    return bool((os.environ.get('LIBREOFFICE_ONLINE_URL') or os.environ.get('COLLABORA_URL') or '').strip())


def collabora_base_url() -> str:
    return (os.environ.get('LIBREOFFICE_ONLINE_URL') or os.environ.get('COLLABORA_URL') or '').strip().rstrip('/')


def wopi_secret() -> str:
    return (
        os.environ.get('CASEPM_WOPI_SECRET')
        or os.environ.get('SECRET_KEY')
        or 'casepm-wopi-dev-secret'
    )


def public_app_base_url(req: Request | None = None) -> str:
    explicit = (os.environ.get('CASEPM_PUBLIC_URL') or os.environ.get('PUBLIC_URL') or '').strip().rstrip('/')
    if explicit:
        return explicit
    req = req or request
    return (req.url_root or '').rstrip('/')


def editor_kind_to_ext(kind: str) -> str:
    return 'xlsx' if (kind or '').strip() == 'sheet' else 'docx'


def extension_for_document(doc) -> str:
    from document_persistence import _editor_kind_for

    kind = _editor_kind_for(doc) or 'doc'
    name = (doc.original_filename or doc.filename or doc.name or '').lower()
    if '.' in name:
        ext = name.rsplit('.', 1)[-1]
        if ext in ('xlsx', 'xls', 'csv', 'ods'):
            return 'xlsx' if ext != 'ods' else 'ods'
        if ext in ('docx', 'doc', 'odt', 'rtf', 'txt', 'html', 'htm'):
            return 'docx' if ext not in ('odt',) else 'odt'
    return editor_kind_to_ext(kind)


def collabora_app_name(ext: str) -> str:
    if ext in ('xlsx', 'xls', 'csv', 'ods'):
        return 'calc'
    return 'writer'


def issue_wopi_token(*, doc_id: int, user_id: int, write: bool = True, ttl_seconds: int = 8 * 3600) -> str:
    now = int(time.time())
    payload = {
        'doc_id': int(doc_id),
        'user_id': int(user_id),
        'write': bool(write),
        'iat': now,
        'exp': now + int(ttl_seconds),
    }
    return jwt.encode(payload, wopi_secret(), algorithm='HS256')


def verify_wopi_token(token: str) -> dict[str, Any] | None:
    if not token:
        return None
    try:
        payload = jwt.decode(token, wopi_secret(), algorithms=['HS256'])
        if 'doc_id' not in payload or 'user_id' not in payload:
            return None
        return payload
    except jwt.PyJWTError:
        return None


def _blank_xlsx_bytes() -> bytes:
    wb = Workbook()
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _blank_docx_bytes() -> bytes:
    try:
        from docx import Document as DocxDocument
    except ImportError:
        return _minimal_docx_bytes()
    doc = DocxDocument()
    doc.add_paragraph('')
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _minimal_docx_bytes() -> bytes:
    """Tiny valid DOCX when python-docx is unavailable."""
    import zipfile

    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""
    rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""
    document = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body><w:p><w:r><w:t></w:t></w:r></w:p></w:body>
</w:document>"""
    buf = BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        zf.writestr('[Content_Types].xml', content_types)
        zf.writestr('_rels/.rels', rels)
        zf.writestr('word/document.xml', document)
    return buf.getvalue()


def blank_office_bytes(ext: str) -> bytes:
    if ext in ('xlsx', 'xls', 'csv'):
        return _blank_xlsx_bytes()
    return _blank_docx_bytes()


def ensure_office_file_on_disk(doc, *, upload_root: str, User=None) -> str:
    """Ensure the document has a real Office file on disk; return absolute path."""
    from document_integration import guess_mime
    from document_persistence import document_folder

    directory = document_folder(upload_root, int(doc.project_id))
    os.makedirs(directory, exist_ok=True)
    path = os.path.join(directory, doc.filename or '')
    ext = extension_for_document(doc)

    if doc.filename and doc.filename != 'pending' and os.path.isfile(path):
        current_ext = (doc.original_filename or doc.filename or '').lower().rsplit('.', 1)[-1]
        if current_ext in ('xlsx', 'xls', 'csv', 'ods', 'docx', 'doc', 'odt'):
            return path

    raw = blank_office_bytes(ext)
    stamp = time.strftime('%Y%m%d%H%M%S')
    safe_name = (doc.name or 'document')[:80].replace('/', '-')
    stored = f'{stamp}_{safe_name}.{ext}'
    full = os.path.join(directory, stored)
    with open(full, 'wb') as fh:
        fh.write(raw)
    if doc.filename and doc.filename not in ('', 'pending'):
        old = os.path.join(directory, doc.filename)
        if os.path.isfile(old) and old != full:
            try:
                os.remove(old)
            except OSError:
                pass
    doc.filename = stored
    doc.original_filename = f'{doc.name or "document"}.{ext}'
    doc.file_size = len(raw)
    doc.mime_type = guess_mime(f'file.{ext}')
    doc.editor_content = None
    return full


def _fetch_discovery_actions(base_url: str) -> dict[str, str]:
    import urllib.request

    now = time.time()
    if _DISCOVERY_CACHE['actions'] and now - _DISCOVERY_CACHE['fetched_at'] < _DISCOVERY_TTL:
        return _DISCOVERY_CACHE['actions']

    url = f'{base_url.rstrip("/")}/hosting/discovery'
    with urllib.request.urlopen(url, timeout=8) as resp:
        xml_text = resp.read().decode('utf-8', errors='replace')
    root = ET.fromstring(xml_text)
    actions: dict[str, str] = {}
    for action in root.findall('.//{*}action'):
        name = (action.attrib.get('name') or '').strip()
        ext = (action.attrib.get('ext') or '').strip().lower()
        urlsrc = (action.attrib.get('urlsrc') or '').strip()
        if name == 'edit' and ext and urlsrc:
            actions[ext] = urlsrc
    _DISCOVERY_CACHE['actions'] = actions
    _DISCOVERY_CACHE['fetched_at'] = now
    return actions


def build_collabora_editor_url(*, wopi_src: str, ext: str) -> str:
    base = collabora_base_url()
    if not base:
        raise RuntimeError('LibreOffice Online URL is not configured.')
    actions = _fetch_discovery_actions(base)
    urlsrc = actions.get(ext.lower())
    if not urlsrc:
        app = collabora_app_name(ext)
        for key, val in actions.items():
            if collabora_app_name(key) == app:
                urlsrc = val
                break
    if not urlsrc:
        raise RuntimeError(f'Collabora discovery has no editor action for .{ext} files.')

    parsed = urlparse(urlsrc)
    query = dict(p.split('=', 1) for p in parsed.query.split('&') if '=' in p)
    query['WOPISrc'] = wopi_src
    query.setdefault('closebutton', '1')
    new_query = urlencode(query, quote_via=quote)
    return urlunparse(parsed._replace(query=new_query))


def build_wopi_src(*, doc_id: int, access_token: str, req: Request | None = None) -> str:
    base = public_app_base_url(req)
    return f'{base}/wopi/files/{int(doc_id)}?access_token={quote(access_token, safe="")}'


def wopi_check_file_info(doc, user, *, file_path: str, write: bool) -> dict[str, Any]:
    size = os.path.getsize(file_path) if os.path.isfile(file_path) else 0
    base_name = doc.original_filename or doc.filename or f'document.{extension_for_document(doc)}'
    friendly = (
        getattr(user, 'full_name', None)
        or getattr(user, 'name', None)
        or getattr(user, 'email', None)
        or f'User {getattr(user, "id", "")}'
    )
    return {
        'BaseFileName': base_name,
        'Size': size,
        'OwnerId': str(getattr(user, 'id', '0')),
        'UserId': str(getattr(user, 'id', '0')),
        'UserFriendlyName': str(friendly),
        'UserCanWrite': bool(write),
        'UserCanNotWriteRelative': True,
        'SupportsUpdate': True,
        'SupportsLocks': False,
        'LastModifiedTime': (doc.updated_at.isoformat() + 'Z') if getattr(doc, 'updated_at', None) else '',
    }
